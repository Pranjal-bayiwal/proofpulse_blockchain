"""File inspection: type sniffing, text extraction, entropy, hidden characters, metadata dates."""
import io
import math
import re
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, timezone

TEXT_EXT = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".htm", ".py", ".js", ".ts", ".java", ".c", ".cpp",
            ".h", ".sol", ".sql", ".yaml", ".yml", ".tex", ".rst", ".log", ".ini", ".css", ".ipynb"}
SIGNATURES = [
    (b"%PDF", "pdf"),
    (b"\x89PNG\r\n\x1a\n", "png"),
    (b"\xff\xd8\xff", "jpeg"),
    (b"GIF87a", "gif"),
    (b"GIF89a", "gif"),
    (b"PK\x03\x04", "zip"),
    (b"MZ", "exe"),
    (b"\x7fELF", "elf"),
    (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", "ole"),
    (b"RIFF", "riff"),
    (b"7z\xbc\xaf\x27\x1c", "7z"),
    (b"Rar!", "rar"),
]
# declared extension -> acceptable sniffed families
EXT_FAMILY = {
    ".pdf": {"pdf"}, ".png": {"png"}, ".jpg": {"jpeg"}, ".jpeg": {"jpeg"}, ".gif": {"gif"},
    ".docx": {"zip"}, ".xlsx": {"zip"}, ".pptx": {"zip"}, ".zip": {"zip"}, ".odt": {"zip"},
    ".doc": {"ole"}, ".xls": {"ole"}, ".ppt": {"ole"},
    ".exe": {"exe"}, ".dll": {"exe"}, ".7z": {"7z"}, ".rar": {"rar"},
}
RISKY_EXT = {".exe", ".dll", ".bat", ".cmd", ".com", ".scr", ".msi", ".js", ".vbs", ".ps1", ".jar", ".sh", ".apk"}
MACRO_EXT = {".docm", ".xlsm", ".pptm"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}
HIDDEN_CHARS = {"​", "‌", "‍", "⁠", "﻿", "‪", "‫", "‬", "‭", "‮",
                "⁦", "⁧", "⁨", "⁩"}
MAX_TEXT_CHARS = 400_000


@dataclass
class Inspection:
    ext: str
    sniffed: str | None
    entropy: float
    text: str = ""
    text_extracted: bool = False
    encrypted_pdf: bool = False
    hidden_chars: int = 0
    has_macros: bool = False
    meta_dates: list[tuple[str, datetime]] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


def ext_of(filename: str) -> str:
    m = re.search(r"(\.[A-Za-z0-9]{1,8})$", filename or "")
    return m.group(1).lower() if m else ""


def sniff(data: bytes) -> str | None:
    head = data[:16]
    for sig, name in SIGNATURES:
        if head.startswith(sig):
            return name
    return None


def shannon_entropy(data: bytes) -> float:
    if not data:
        return 0.0
    sample = data[:1_000_000]
    counts = Counter(sample)
    n = len(sample)
    return -sum((c / n) * math.log2(c / n) for c in counts.values())


def _aware(dt: datetime) -> datetime:
    return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)


def _extract_pdf(data: bytes, ins: Inspection) -> None:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            ins.encrypted_pdf = True
            return
        parts = []
        for page in reader.pages[:200]:
            parts.append(page.extract_text() or "")
            if sum(map(len, parts)) > MAX_TEXT_CHARS:
                break
        ins.text = "\n".join(parts)
        ins.text_extracted = True
        meta = reader.metadata
        if meta:
            for key, label in (("creation_date", "PDF created"), ("modification_date", "PDF modified")):
                try:
                    dt = getattr(meta, key)
                    if dt:
                        ins.meta_dates.append((label, _aware(dt)))
                except Exception:
                    pass
    except Exception as exc:
        ins.notes.append(f"PDF could not be parsed ({type(exc).__name__})")


def _extract_docx(data: bytes, ins: Inspection) -> None:
    try:
        from docx import Document

        doc = Document(io.BytesIO(data))
        ins.text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                ins.text += "\n" + " ".join(c.text for c in row.cells)
        ins.text_extracted = True
        cp = doc.core_properties
        for label, dt in (("DOCX created", cp.created), ("DOCX modified", cp.modified)):
            if dt:
                ins.meta_dates.append((label, _aware(dt)))
    except Exception as exc:
        ins.notes.append(f"DOCX could not be parsed ({type(exc).__name__})")


def inspect_file(filename: str, data: bytes) -> Inspection:
    ext = ext_of(filename)
    ins = Inspection(ext=ext, sniffed=sniff(data), entropy=shannon_entropy(data))

    if ins.sniffed == "zip":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = z.namelist()
                ins.has_macros = any(n.lower().endswith("vbaproject.bin") for n in names)
        except zipfile.BadZipFile:
            ins.notes.append("Archive structure is corrupt")

    if ext == ".pdf" and ins.sniffed == "pdf":
        _extract_pdf(data, ins)
    elif ext == ".docx" and ins.sniffed == "zip":
        _extract_docx(data, ins)
    elif ext in TEXT_EXT or (ins.sniffed is None and ext not in IMAGE_EXT and ins.entropy < 6.5):
        try:
            ins.text = data[: MAX_TEXT_CHARS * 2].decode("utf-8", errors="replace")
            ins.text_extracted = True
        except Exception:
            pass

    if ins.text:
        ins.text = ins.text[:MAX_TEXT_CHARS]
        ins.hidden_chars = sum(1 for ch in ins.text if ch in HIDDEN_CHARS)
        # also count other invisible format characters (category Cf)
        ins.hidden_chars += sum(1 for ch in ins.text if unicodedata.category(ch) == "Cf" and ch not in HIDDEN_CHARS)
    return ins


def normalise_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).lower()
    text = "".join(ch for ch in text if ch not in HIDDEN_CHARS)
    return re.sub(r"[^\w]+", " ", text, flags=re.UNICODE).strip()
